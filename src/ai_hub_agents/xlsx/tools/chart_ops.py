"""图表工具。"""

from __future__ import annotations

from typing import Annotated, Optional

from langchain_core.tools import ToolException, tool
from langgraph.prebuilt import InjectedState
from openpyxl import load_workbook

from ._helpers import get_output_file, with_file_lock

_CHART_CLASSES = {
    "bar": "BarChart",
    "column": "BarChart",
    "line": "LineChart",
    "pie": "PieChart",
    "area": "AreaChart",
    "scatter": "ScatterChart",
    "doughnut": "DoughnutChart",
}


@tool
@with_file_lock
def create_chart(
    state: Annotated[dict, InjectedState],
    sheet_name: str,
    chart_type: str,
    data_range: str,
    target_cell: str,
    title: Optional[str] = None,
    x_axis_title: Optional[str] = None,
    y_axis_title: Optional[str] = None,
    width: float = 15.0,
    height: float = 10.0,
    categories_range: Optional[str] = None,
) -> str:
    """在输出文件中创建图表。

    Args:
        sheet_name: 工作表名称
        chart_type: 图表类型，支持 "bar", "column", "line", "pie", "area", "scatter", "doughnut"
        data_range: 数据范围如 "B1:D10"（含标题行）
        target_cell: 图表放置位置，如 "F2"
        title: 图表标题
        x_axis_title: X 轴标题
        y_axis_title: Y 轴标题
        width: 图表宽度（厘米），默认 15.0
        height: 图表高度（厘米），默认 10.0
        categories_range: 类别标签范围如 "A2:A10"
    """
    from openpyxl.chart import Reference
    from openpyxl.utils import range_boundaries

    cls_name = _CHART_CLASSES.get(chart_type)
    if not cls_name:
        raise ToolException(
            f"不支持的图表类型: {chart_type}，可选: {', '.join(_CHART_CLASSES)}"
        )

    import openpyxl.chart as chart_mod

    ChartClass = getattr(chart_mod, cls_name)

    output = get_output_file(state)
    try:
        wb = load_workbook(output)
        ws = wb[sheet_name]

        chart = ChartClass()
        if chart_type == "column":
            chart.type = "col"

        if title:
            chart.title = title
        if x_axis_title and hasattr(chart, "x_axis"):
            chart.x_axis.title = x_axis_title
        if y_axis_title and hasattr(chart, "y_axis"):
            chart.y_axis.title = y_axis_title
        chart.width = width
        chart.height = height

        min_col, min_row, max_col, max_row = range_boundaries(data_range)
        data_ref = Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
        chart.add_data(data_ref, titles_from_data=True)

        if categories_range:
            c_min_col, c_min_row, c_max_col, c_max_row = range_boundaries(categories_range)
            cats = Reference(ws, min_col=c_min_col, min_row=c_min_row, max_col=c_max_col, max_row=c_max_row)
            chart.set_categories(cats)

        ws.add_chart(chart, target_cell)
        wb.save(output)
        return f"已在 {sheet_name}!{target_cell} 创建 {chart_type} 图表"
    except KeyError:
        raise ToolException(f"工作表 '{sheet_name}' 不存在")
    except ToolException:
        raise
    except Exception as e:
        raise ToolException(f"create_chart 失败: {e}") from e
